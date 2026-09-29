"""端到端集成测试：用命令行入口跑完整流程。

这些测试**不碰用户的真实资料目录**，全部在 pytest 的 tmp_path 里造数据。
PDF 材料识别那一层会被替换成桩（真的造一个带文字层、可被 pdfplumber 解析的
PDF 需要额外依赖，收益不大），材料规则本身由 ``test_material.py`` 覆盖。
"""

from __future__ import annotations

import os

import openpyxl
import pytest

from speedup import cli
from speedup.config import COL_MATERIAL, COL_NAME, COL_NOTE, DATA_START_ROW
from speedup.excel.reader import extract_names
from speedup.material import MaterialLookup
from speedup.techreq import CraftLookup


@pytest.fixture
def fake_materials(monkeypatch):
    """让「同目录同名 PDF」固定返回指定材料。

    用法：``fake_materials({"关节支撑板": "7075-T6"})``
    """

    def _apply(mapping: dict[str, str]):
        def fake_lookup(model_path, pdf_index):
            stem = os.path.splitext(os.path.basename(model_path))[0]
            if stem not in mapping:
                return MaterialLookup(), None
            return MaterialLookup(material=mapping[stem], where="标题栏"), f"{stem}.pdf"

        monkeypatch.setattr("speedup.tools.gbbuild.lookup_material_for", fake_lookup)

    return _apply


@pytest.fixture
def fake_crafts(monkeypatch):
    """让「读图纸技术要求栏」固定返回指定工艺。

    用法：``fake_crafts({"关节支撑板": "阳极氧化本色"})``。
    值为 ``None`` 表示**图纸里没写工艺**（和「读不了」不是一回事）。
    """

    def _apply(mapping: dict[str, str | None]):
        def fake_craft(pdf_path, **_kwargs):
            stem = os.path.splitext(os.path.basename(str(pdf_path)))[0]
            if stem not in mapping:
                return CraftLookup(where="技术要求：")
            text = mapping[stem]
            if text is None:
                return CraftLookup(where="技术要求：")
            return CraftLookup(text=text, raw=f"5、{text}；", where="技术要求：")

        monkeypatch.setattr("speedup.tools.gbbuild.find_craft_in_pdf", fake_craft)

    return _apply


class TestGbbuildEndToEnd:
    def test_generates_list_from_step_files(self, tmp_path, make_files, fake_materials, capsys):
        folder = tmp_path / "0916"
        make_files(folder, ["左臂末端基座.STEP", "垫片.STEP", "关节支撑板.stp"])
        fake_materials({"左臂末端基座": "7075-T6", "关节支撑板": "7075-T6"})

        assert cli.main(["gbbuild", str(folder), "--no-open"]) == 0

        destination = folder / "0916.xlsx"
        assert destination.exists()
        names = list(extract_names(destination).names)
        # 顺序 = 模型文件名的排序（按码位），换零件名会连带改变顺序
        assert names == ["关节支撑板", "垫片", "左臂末端基座"]

        out = capsys.readouterr().out
        assert "共写入 3 个零件" in out
        assert "垫片" in out  # 会出现在「没有同名图纸」的警告里

    def test_material_lands_in_the_right_column(self, tmp_path, make_files, fake_materials):
        folder = tmp_path / "0916"
        make_files(folder, ["关节支撑板.STEP"])
        fake_materials({"关节支撑板": "7075-T6"})

        cli.main(["gbbuild", str(folder), "--no-open"])

        worksheet = openpyxl.load_workbook(folder / "0916.xlsx").active
        assert worksheet.cell(row=DATA_START_ROW, column=COL_NAME).value == "关节支撑板"
        assert worksheet.cell(row=DATA_START_ROW, column=COL_MATERIAL).value == "7075-T6"

    def test_custom_output_filename(self, tmp_path, make_files):
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        assert cli.main(["gbbuild", str(folder), "--out", "采购清单.xlsx", "--no-open"]) == 0
        assert (folder / "采购清单.xlsx").exists()

    def test_scans_subdirectories(self, tmp_path, make_files):
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        make_files(folder / "子目录", ["法兰盘.STEP"])
        cli.main(["gbbuild", str(folder), "--no-open"])
        names = list(extract_names(folder / "0916.xlsx").names)
        assert names == ["垫片", "法兰盘"]

    def test_ignores_non_model_files(self, tmp_path, make_files):
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP", "0916.xlsx", "J7关节_导出清单.csv", "说明.txt"])
        cli.main(["gbbuild", str(folder), "--no-open"])
        names = list(extract_names(folder / "0916.xlsx").names)
        assert names == ["垫片"]

    def test_empty_folder_returns_2_with_message(self, tmp_path, capsys):
        folder = tmp_path / "空的"
        folder.mkdir()
        assert cli.main(["gbbuild", str(folder), "--no-open"]) == 2
        assert "没有找到" in capsys.readouterr().out

    def test_second_run_reports_overwrite(self, tmp_path, make_files, capsys):
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        cli.main(["gbbuild", str(folder), "--no-open"])
        capsys.readouterr()
        cli.main(["gbbuild", str(folder), "--no-open"])
        assert "覆盖旧文件" in capsys.readouterr().out

    def test_craft_lands_in_the_note_column(
        self, tmp_path, make_files, fake_materials, fake_crafts
    ):
        """图纸技术要求栏里的表面处理要求，要落到备注列。"""
        folder = tmp_path / "0916"
        make_files(folder, ["关节支撑板.STEP"])
        fake_materials({"关节支撑板": "7075-T6"})
        fake_crafts({"关节支撑板": "阳极氧化本色"})

        cli.main(["gbbuild", str(folder), "--no-open"])

        worksheet = openpyxl.load_workbook(folder / "0916.xlsx").active
        assert (
            worksheet.cell(row=DATA_START_ROW, column=COL_NOTE).value
            == "阳极氧化本色"
        )

    def test_part_without_craft_line_is_reported(
        self, tmp_path, make_files, fake_materials, fake_crafts, capsys
    ):
        """有图纸、但技术要求栏里没写工艺 —— 要在汇总里点名。

        不点名的话，备注列那个空格没人会当回事，等加工厂问起来才发现漏了。
        """
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        fake_materials({"垫片": "Q235"})
        fake_crafts({"垫片": None})

        cli.main(["gbbuild", str(folder), "--no-open"])

        out = capsys.readouterr().out
        assert "表面处理" in out
        assert "垫片" in out

    def test_read_error_is_not_reported_as_missing(
        self, tmp_path, make_files, fake_materials, capsys, monkeypatch
    ):
        """「读不了」和「图纸没写」必须分开报（R26 的最后一条）。"""
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        fake_materials({"垫片": "Q235"})
        monkeypatch.setattr(
            "speedup.tools.gbbuild.find_craft_in_pdf",
            lambda _path, **_kw: CraftLookup(error="读取失败(假的)"),
        )

        cli.main(["gbbuild", str(folder), "--no-open"])

        out = capsys.readouterr().out
        assert "没有写表面处理" not in out
        assert "读取失败" in out

    def test_part_without_pdf_has_empty_note(
        self, tmp_path, make_files, fake_materials, fake_crafts
    ):
        """没有同名图纸时不该去读工艺，备注列也该是空的。"""
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        fake_materials({})
        fake_crafts({})

        cli.main(["gbbuild", str(folder), "--no-open"])

        worksheet = openpyxl.load_workbook(folder / "0916.xlsx").active
        assert worksheet.cell(row=DATA_START_ROW, column=COL_NOTE).value is None

    def test_missing_pdf_lists_the_folder_pdfs(
        self, tmp_path, make_files, fake_materials, capsys
    ):
        """没匹配到图纸时，把文件夹里实际有哪些 PDF 打出来。

        这一条是给「明明有图纸，为什么说没有」准备的 —— 一眼就能分清
        「图纸真不在这个文件夹」和「图纸在，只是文件名对不上」。
        """
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP", "别的东西.pdf"])
        fake_materials({})

        cli.main(["gbbuild", str(folder), "--no-open"])

        out = capsys.readouterr().out
        assert "别的东西.pdf" in out

    def test_missing_pdf_when_folder_has_none(
        self, tmp_path, make_files, fake_materials, capsys
    ):
        folder = tmp_path / "0916"
        make_files(folder, ["垫片.STEP"])
        fake_materials({})

        cli.main(["gbbuild", str(folder), "--no-open"])

        out = capsys.readouterr().out
        assert "一张 PDF 都没有" in out


class TestGbcopyEndToEnd:
    def _make_list(self, folder, names):
        """造一份真实的采购清单（用 gbbuild 的格式），而不是手搓一个假的。"""
        from speedup.excel.writer import generate_supplement_list
        from speedup.models import SupplementEntry

        folder.mkdir(parents=True, exist_ok=True)
        destination, _existed = generate_supplement_list(
            folder, [SupplementEntry(name=name) for name in names], filename="采购清单.xlsx"
        )
        return destination

    def test_copies_matching_files(self, tmp_path, make_files):
        target = tmp_path / "清单"
        listing = self._make_list(target, ["关节支撑板", "垫片"])
        source = tmp_path / "source"
        make_files(source, ["关节支撑板V1.1.STEP", "关节支撑板.pdf", "垫片.STEP"])

        assert cli.main(["gbcopy", str(listing), str(source), "--no-open"]) == 0

        assert (target / "关节支撑板V1.1.STEP").exists()
        assert (target / "关节支撑板.pdf").exists()
        assert (target / "垫片.STEP").exists()

    def test_dry_run_copies_nothing(self, tmp_path, make_files, capsys):
        target = tmp_path / "清单"
        listing = self._make_list(target, ["垫片"])
        source = tmp_path / "source"
        make_files(source, ["垫片.STEP"])

        assert cli.main(["gbcopy", str(listing), str(source), "--dry-run"]) == 0

        assert not (target / "垫片.STEP").exists()
        assert "没有真的拷贝" in capsys.readouterr().out

    def test_does_not_overwrite_existing_file(self, tmp_path, make_files):
        target = tmp_path / "清单"
        listing = self._make_list(target, ["垫片"])
        (existing,) = make_files(target, ["垫片.STEP"])
        existing.write_text("人工改过，别覆盖", encoding="utf-8")

        source = tmp_path / "source"
        (source_file,) = make_files(source, ["垫片.STEP"])
        source_file.write_text("新的", encoding="utf-8")

        cli.main(["gbcopy", str(listing), str(source), "--no-open"])

        assert existing.read_text(encoding="utf-8") == "人工改过，别覆盖"

    def test_strict_mode_skips_version_suffix(self, tmp_path, make_files, capsys):
        target = tmp_path / "清单"
        listing = self._make_list(target, ["关节支撑板"])
        source = tmp_path / "source"
        make_files(source, ["关节支撑板V1.1.STEP"])

        cli.main(["gbcopy", str(listing), str(source), "--min-score", "4", "--no-open"])

        assert not (target / "关节支撑板V1.1.STEP").exists()
        assert "未找到" in capsys.readouterr().out

    def test_reports_missing_files(self, tmp_path, make_files, capsys):
        target = tmp_path / "清单"
        listing = self._make_list(target, ["压根没有的零件"])
        source = tmp_path / "source"
        make_files(source, ["垫片.STEP"])

        cli.main(["gbcopy", str(listing), str(source), "--no-open"])
        out = capsys.readouterr().out
        assert "完全没找到 1 个名称" in out

    def test_missing_listing_file_returns_2(self, tmp_path, capsys):
        source = tmp_path / "source"
        source.mkdir()
        code = cli.main(["gbcopy", str(tmp_path / "没有.xlsx"), str(source), "--no-open"])
        assert code == 2
        assert "[失败]" in capsys.readouterr().out

    def test_empty_source_folder_returns_2(self, tmp_path, make_files, capsys):
        target = tmp_path / "清单"
        listing = self._make_list(target, ["垫片"])
        source = tmp_path / "source"
        source.mkdir()

        assert cli.main(["gbcopy", str(listing), str(source), "--no-open"]) == 2
        assert "没有找到任何" in capsys.readouterr().out

    def test_legacy_xls_listing_is_rejected(self, tmp_path, capsys):
        legacy = tmp_path / "老清单.xls"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0")
        source = tmp_path / "source"
        source.mkdir()

        assert cli.main(["gbcopy", str(legacy), str(source), "--no-open"]) == 2
        assert "xls" in capsys.readouterr().out


class TestMultiSheetListing:
    """回归测试：工作簿里有好几张带「名称」列的表时，只能认一张。

    背景（`AGENTS.md` P12）：用户那份 `0825示例批次.xlsx` 有「采购清单」和
    「备料参考」两张表。旧实现把两张表的名**合并**成 47 个，
    于是参考页的 26 个零件也被拷了出去 —— 用户是在目标文件夹里看到
    「怎么多了一堆图纸」才发现不对的。
    """

    @staticmethod
    def _make_multi_sheet_list(folder, groups):
        """按 ``{表名: [名称, ...]}`` 造一份多表清单。第一张表在前。"""
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / "清单.xlsx"
        workbook = openpyxl.Workbook()
        for index, (sheet_title, names) in enumerate(groups.items()):
            sheet = workbook.active if index == 0 else workbook.create_sheet(sheet_title)
            sheet.title = sheet_title
            sheet.append(["序号", "名称", "材料"])
            for order, name in enumerate(names, 1):
                sheet.append([order, name, ""])
        workbook.save(path)
        return path

    @pytest.fixture
    def scene(self, tmp_path, make_files):
        """还原用户的现场：主表 1 个零件，参考表 3 个零件，库里有全部 4 个的模型。"""
        target = tmp_path / "0928"
        listing = self._make_multi_sheet_list(
            target,
            {
                "采购清单": ["座板"],
                "备料参考": ["NECK", "HEAD_BACK", "隔离环"],
            },
        )
        source = tmp_path / "模型库"
        make_files(
            source,
            ["座板_V1.2.STEP", "座板_V1.2.pdf", "NECK.STEP", "HEAD_BACK.STEP", "隔离环.STEP"],
        )
        return target, listing, source

    def test_only_the_main_sheet_gets_copied(self, scene):
        target, listing, source = scene
        assert cli.main(["gbcopy", str(listing), str(source), "--no-open"]) == 0

        assert (target / "座板_V1.2.STEP").exists()
        assert (target / "座板_V1.2.pdf").exists()
        for leaked in ("NECK.STEP", "HEAD_BACK.STEP", "隔离环.STEP"):
            assert not (target / leaked).exists(), f"参考表的 {leaked} 又被拷出来了"

    def test_warns_about_the_ignored_sheet(self, scene, capsys):
        target, listing, source = scene
        cli.main(["gbcopy", str(listing), str(source), "--no-open"])
        out = capsys.readouterr().out
        assert "[警告]" in out
        assert "备料参考" in out  # 说清忽略了哪张表
        assert "--sheet" in out  # 给出路
        # 提醒要出现在名称列表之前，否则一屏名字刷过去根本反应不过来
        assert out.index("备料参考") < out.index("个零件：")

    def test_sheet_flag_overrides_the_choice(self, scene):
        target, listing, source = scene
        assert (
            cli.main(
                ["gbcopy", str(listing), str(source), "--sheet", "备料参考", "--no-open"]
            )
            == 0
        )
        assert (target / "NECK.STEP").exists()
        assert not (target / "座板_V1.2.STEP").exists(), "指定了表就不该再读主表"

    def test_sheet_flag_silences_the_warning(self, scene, capsys):
        target, listing, source = scene
        cli.main(["gbcopy", str(listing), str(source), "--sheet", "采购清单", "--no-open"])
        assert "[警告]" not in capsys.readouterr().out, "用户已经明确指定了，不用再唠叨"

    def test_unknown_sheet_name_returns_2(self, scene, capsys):
        target, listing, source = scene
        assert (
            cli.main(["gbcopy", str(listing), str(source), "--sheet", "没有这张", "--no-open"]) == 2
        )
        out = capsys.readouterr().out
        assert "[失败]" in out
        assert "采购清单" in out  # 报错要告诉人有哪些表可选


class TestFullWorkflow:
    def test_build_then_copy(self, tmp_path, make_files, fake_materials):
        """完整链路：模型目录 → 生成清单 → 换台机器按清单回收文件。"""
        source = tmp_path / "零件库"
        make_files(source, ["垫片.STEP", "法兰盘.stp"])
        fake_materials({"垫片": "Q235"})

        # 1) 生成清单
        assert cli.main(["gbbuild", str(source), "--no-open"]) == 0
        listing = source / "零件库.xlsx"
        assert listing.exists()

        # 2) 库里多出一个清单里没有的零件 —— 它不该被拷走
        make_files(source, ["多余零件.STEP"])

        # 3) 把清单挪到另一个文件夹，模拟「清单在手上，文件在库里」
        work = tmp_path / "外发"
        work.mkdir()
        moved = work / listing.name
        listing.replace(moved)

        assert cli.main(["gbcopy", str(moved), str(source), "--no-open"]) == 0
        assert (work / "垫片.STEP").exists()
        assert (work / "法兰盘.stp").exists()
        assert not (work / "多余零件.STEP").exists()
