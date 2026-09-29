"""采购清单生成测试：表头格式、数据落位、命名、覆盖标记。"""

from __future__ import annotations

import os

import openpyxl
import pytest

from speedup.config import (
    CENTERED_COLUMNS,
    COL_MATERIAL,
    COL_NAME,
    COL_NO,
    COL_NOTE,
    COL_QTY,
    DATA_ROW_HEIGHT,
    DATA_START_ROW,
    HEADER_FILL,
    HEADER_FONT_COLOR,
    HEADER_ROW,
    HEADER_ROW_HEIGHT,
    MAX_ROWS,
    MIN_DATA_ROWS,
    SHEET_TITLE,
    SUPPLEMENT_HEADERS,
    SUPPLEMENT_WIDTHS,
)
from speedup.excel.reader import extract_names
from speedup.excel.writer import (
    build_supplement_workbook,
    fill_supplement_sheet,
    generate_supplement_list,
)
from speedup.models import SupplementEntry


class TestBuildWorkbook:
    def test_sheet_title_and_headers(self):
        _workbook, worksheet = build_supplement_workbook(0)
        assert worksheet.title == SHEET_TITLE
        actual = [worksheet.cell(row=HEADER_ROW, column=i).value for i in range(1, len(SUPPLEMENT_HEADERS) + 1)]
        assert tuple(actual) == SUPPLEMENT_HEADERS

    def test_header_is_styled(self):
        _workbook, worksheet = build_supplement_workbook(0)
        cell = worksheet.cell(row=HEADER_ROW, column=1)
        assert cell.font.bold is True
        assert cell.font.color.rgb.endswith(HEADER_FONT_COLOR)
        assert cell.fill.fgColor.rgb.endswith(HEADER_FILL)
        assert cell.fill.patternType == "solid"
        assert worksheet.row_dimensions[HEADER_ROW].height == HEADER_ROW_HEIGHT

    def test_column_widths_applied(self):
        from openpyxl.utils import get_column_letter

        _workbook, worksheet = build_supplement_workbook(0)
        for index, width in enumerate(SUPPLEMENT_WIDTHS, 1):
            assert worksheet.column_dimensions[get_column_letter(index)].width == width

    @pytest.mark.parametrize("column", CENTERED_COLUMNS)
    def test_short_number_columns_are_centered(self, column):
        """序号、数量是短数字，居中才像一张表 —— 左对齐会贴着格线看着散。"""
        _workbook, worksheet = build_supplement_workbook(0)
        for row in range(DATA_START_ROW, MIN_DATA_ROWS + 1):
            assert worksheet.cell(row=row, column=column).alignment.horizontal == "center"

    def test_text_columns_stay_left_aligned(self):
        """名称、材料、备注是长短不一的文字，居中反而难扫读。"""
        _workbook, worksheet = build_supplement_workbook(0)
        text_columns = [
            index
            for index in range(1, len(SUPPLEMENT_HEADERS) + 1)
            if index not in CENTERED_COLUMNS
        ]
        assert text_columns, "至少要有一列是左对齐的，否则这条测试就没意义了"
        for column in text_columns:
            for row in range(DATA_START_ROW, MIN_DATA_ROWS + 1):
                alignment = worksheet.cell(row=row, column=column).alignment
                assert alignment.horizontal in (None, "left")

    def test_center_columns_are_vertically_centered_too(self):
        """居中的同时也要垂直居中，别只改一半。"""
        _workbook, worksheet = build_supplement_workbook(0)
        for column in CENTERED_COLUMNS:
            alignment = worksheet.cell(row=DATA_START_ROW, column=column).alignment
            assert alignment.vertical == "center"

    def test_empty_table_still_gets_minimum_rows(self):
        _workbook, worksheet = build_supplement_workbook(0)
        for row in range(DATA_START_ROW, MIN_DATA_ROWS + 1):
            assert worksheet.row_dimensions[row].height == DATA_ROW_HEIGHT
            assert worksheet.cell(row=row, column=1).border.left.style == "thin"

    def test_formats_exactly_enough_rows(self):
        _workbook, worksheet = build_supplement_workbook(3)
        assert worksheet.row_dimensions[DATA_START_ROW + 2].height == DATA_ROW_HEIGHT

    def test_does_not_format_beyond_max_rows(self):
        _workbook, worksheet = build_supplement_workbook(MAX_ROWS + 50)
        assert worksheet.row_dimensions[MAX_ROWS].height == DATA_ROW_HEIGHT
        assert worksheet.row_dimensions[MAX_ROWS + 1].height is None


class TestFillSheet:
    def test_writes_name_material_quantity(self):
        _workbook, worksheet = build_supplement_workbook(2)
        written = fill_supplement_sheet(
            worksheet,
            [
                SupplementEntry(name="垫片", material="7075-T6"),
                SupplementEntry(name="法兰盘"),
            ],
        )
        assert written == 2

        first = DATA_START_ROW
        assert worksheet.cell(row=first, column=COL_NO).value == 1
        assert worksheet.cell(row=first, column=COL_NAME).value == "垫片"
        assert worksheet.cell(row=first, column=COL_MATERIAL).value == "7075-T6"
        assert worksheet.cell(row=first, column=COL_QTY).value == 1

    def test_leaves_material_blank_when_unknown(self):
        _workbook, worksheet = build_supplement_workbook(1)
        fill_supplement_sheet(worksheet, [SupplementEntry(name="垫片")])
        assert worksheet.cell(row=DATA_START_ROW, column=COL_MATERIAL).value is None

    def test_writes_note_when_present(self):
        _workbook, worksheet = build_supplement_workbook(1)
        fill_supplement_sheet(worksheet, [SupplementEntry(name="垫片", note="未找到同名图纸")])
        assert worksheet.cell(row=DATA_START_ROW, column=COL_NOTE).value == "未找到同名图纸"

    def test_numbers_are_sequential(self):
        _workbook, worksheet = build_supplement_workbook(3)
        fill_supplement_sheet(worksheet, [SupplementEntry(name=f"件{i}") for i in range(3)])
        numbers = [worksheet.cell(row=DATA_START_ROW + i, column=COL_NO).value for i in range(3)]
        assert numbers == [1, 2, 3]

    def test_stops_at_max_rows(self):
        total = MAX_ROWS + 5
        _workbook, worksheet = build_supplement_workbook(total)
        written = fill_supplement_sheet(worksheet, [SupplementEntry(name=f"件{i}") for i in range(total)])
        assert written == MAX_ROWS - DATA_START_ROW + 1

    def test_empty_entries(self):
        _workbook, worksheet = build_supplement_workbook(0)
        assert fill_supplement_sheet(worksheet, []) == 0


class TestGenerateSupplementList:
    def test_names_file_after_folder(self, tmp_path):
        folder = tmp_path / "0916"
        folder.mkdir()
        destination, existed = generate_supplement_list(folder, [SupplementEntry(name="垫片")])

        assert destination.endswith("0916.xlsx")
        assert existed is False
        assert (folder / "0916.xlsx").exists()

    def test_custom_filename(self, tmp_path):
        destination, _existed = generate_supplement_list(
            tmp_path, [SupplementEntry(name="垫片")], filename="自定义.xlsx"
        )
        assert destination.endswith("自定义.xlsx")

    def test_reports_overwrite(self, tmp_path):
        generate_supplement_list(tmp_path, [SupplementEntry(name="垫片")])
        _destination, existed = generate_supplement_list(tmp_path, [SupplementEntry(name="垫片")])
        assert existed is True

    def test_handles_trailing_separator_in_folder(self, tmp_path):
        folder = tmp_path / "0916"
        folder.mkdir()
        destination, _existed = generate_supplement_list(str(folder) + os.sep, [])
        assert destination.endswith("0916.xlsx")

    def test_output_is_readable_by_our_own_reader(self, tmp_path):
        """闭环测试：生成 → 读回，确认「名称」列真的能被 gbcopy 用上。"""
        folder = tmp_path / "0916"
        folder.mkdir()
        entries = [SupplementEntry(name="垫片"), SupplementEntry(name="法兰盘")]
        generate_supplement_list(folder, entries)

        names = list(extract_names(folder / "0916.xlsx").names)
        assert names == ["垫片", "法兰盘"]

    def test_missing_folder_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            generate_supplement_list(tmp_path / "没有这个目录", [])

    def test_generated_workbook_has_no_stray_default_sheet(self, tmp_path):
        destination, _existed = generate_supplement_list(tmp_path, [SupplementEntry(name="垫片")])
        workbook = openpyxl.load_workbook(destination)
        assert workbook.sheetnames == [SHEET_TITLE]
