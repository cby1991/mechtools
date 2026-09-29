"""Excel 读取测试：两种解析器往返、名称列抽取、异常路径。"""

from __future__ import annotations

import openpyxl
import pytest

from speedup.config import (
    LEGACY_SHEET_TITLE,
    NAME_HEADER_KEYWORD,
    PREFERRED_INPUT_SHEETS,
    SHEET_TITLE,
)
from speedup.errors import InputError, UnsupportedFormatError
from speedup.excel.reader import (
    column_letters_to_index,
    extract_names,
    extract_names_from_sheets,
    read_workbook,
    read_xlsx_builtin,
    read_xlsx_with_openpyxl,
)


def _write_book(path, rows, title="采购清单", extra_sheets=None):
    """造一个 xlsx。``extra_sheets`` 用来加「参考表」这种同表头的干扰项。"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = title
    for row in rows:
        worksheet.append(list(row))
    for sheet_title, sheet_rows in (extra_sheets or {}).items():
        extra = workbook.create_sheet(sheet_title)
        for row in sheet_rows:
            extra.append(list(row))
    workbook.save(path)
    return path


class TestColumnLettersToIndex:
    @pytest.mark.parametrize(
        ("letters", "expected"),
        [("A", 0), ("B", 1), ("C", 2), ("Z", 25), ("AA", 26), ("AB", 27), ("BA", 52)],
    )
    def test_converts(self, letters, expected):
        assert column_letters_to_index(letters) == expected

    def test_case_insensitive(self):
        assert column_letters_to_index("aa") == column_letters_to_index("AA")


class TestBothParsersAgree:
    """openpyxl 和内置解析器必须给出同样的结果，否则回退路径会悄悄改变行为。"""

    def test_simple_sheet(self, tmp_path):
        path = _write_book(tmp_path / "a.xlsx", [["序号", "名称"], [1, "垫片"], [2, "法兰盘"]])
        expected = [["序号", "名称"], ["1", "垫片"], ["2", "法兰盘"]]

        assert read_xlsx_with_openpyxl(path)["采购清单"] == expected
        assert read_xlsx_builtin(path)["采购清单"] == expected

    def test_empty_cells_and_gaps(self, tmp_path):
        path = _write_book(tmp_path / "b.xlsx", [["名称", "", "材料"], ["垫片", None, "7075-T6"]])
        sheets = read_xlsx_builtin(path)
        assert sheets["采购清单"][1][0] == "垫片"
        assert sheets["采购清单"][1][2] == "7075-T6"

    def test_multiple_sheets(self, tmp_path):
        workbook = openpyxl.Workbook()
        first = workbook.active
        first.title = "表一"
        first.append(["名称"])
        first.append(["垫片"])
        second = workbook.create_sheet("表二")
        second.append(["名称"])
        second.append(["法兰盘"])
        path = tmp_path / "c.xlsx"
        workbook.save(path)

        for parser in (read_xlsx_with_openpyxl, read_xlsx_builtin):
            sheets = parser(path)
            assert set(sheets) == {"表一", "表二"}


class TestReadWorkbook:
    def test_reads_xlsx(self, tmp_path):
        path = _write_book(tmp_path / "a.xlsx", [["名称"], ["垫片"]])
        assert "采购清单" in read_workbook(path)

    def test_rejects_legacy_xls(self, tmp_path):
        legacy = tmp_path / "old.xls"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0")
        with pytest.raises(UnsupportedFormatError, match="xls"):
            read_workbook(legacy)

    def test_missing_file(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            read_workbook(tmp_path / "没有.xlsx")

    def test_falls_back_to_builtin_when_openpyxl_raises(self, tmp_path, monkeypatch):
        path = _write_book(tmp_path / "a.xlsx", [["名称"], ["垫片"]])

        def boom(*args, **kwargs):
            raise ValueError("模拟 openpyxl 解析不了这个文件")

        monkeypatch.setattr("speedup.excel.reader.read_xlsx_with_openpyxl", boom)
        assert read_workbook(path)["采购清单"][1][0] == "垫片"


class TestExtractNamesFromSheets:
    def test_basic(self):
        result = extract_names_from_sheets({"S": [["序号", "名称"], [1, "垫片"], [2, "法兰盘"]]})
        assert list(result.names) == ["垫片", "法兰盘"]
        assert result.sheet == "S"
        assert result.column == 2
        assert result.rows == 2
        assert "工作表「S」第 2 列" in result.summary()
        assert "读到 2 行" in result.summary()

    def test_header_can_be_anywhere_in_first_ten_rows(self):
        rows = [[None]] * 5 + [["名称"], ["垫片"]]
        assert list(extract_names_from_sheets({"S": rows}).names) == ["垫片"]

    def test_header_not_found(self):
        result = extract_names_from_sheets({"S": [["序号", "图号"], [1, "a"]]})
        assert result.names == ()
        assert result.sheet is None
        assert f"未找到名为「{NAME_HEADER_KEYWORD}」的列" in result.summary()

    def test_deduplicates_preserving_order(self):
        rows = [["名称"], ["垫片"], ["法兰盘"], ["垫片"]]
        result = extract_names_from_sheets({"S": rows})
        assert list(result.names) == ["垫片", "法兰盘"]
        assert result.rows == 3  # 去重前读到 3 行

    def test_skips_blank_and_none_like_values(self):
        rows = [["名称"], ["垫片"], [None], [""], ["none"], [" 法兰盘 "]]
        assert list(extract_names_from_sheets({"S": rows}).names) == ["垫片", "法兰盘"]

    def test_short_rows_are_skipped(self):
        rows = [["序号", "名称"], [1], [2, "垫片"]]
        assert list(extract_names_from_sheets({"S": rows}).names) == ["垫片"]

    def test_custom_header_keyword(self):
        result = extract_names_from_sheets({"S": [["零件"], ["垫片"]]}, header_keyword="零件")
        assert list(result.names) == ["垫片"]
        assert result.header_keyword == "零件"

    def test_empty_sheets(self):
        result = extract_names_from_sheets({})
        assert result.names == ()
        assert "未找到" in result.summary()


class TestSheetSelection:
    """一份工作簿里可能有好几张都带「名称」列的表，只该认一张。

    回归背景（`AGENTS.md` P12）：`0825示例批次.xlsx` 有「采购清单」和
    「备料参考」两张表，旧实现把两张表的名**合并**成 47 个，
    结果多拷了 26 个文件出去。
    """

    TWO_SHEETS = {
        "采购清单": [["序号", "名称"], [1, "座板"], [2, "臂杆"]],
        "备料参考": [["序号", "名称"], [1, "NECK"], [2, "HEAD_BACK"], [3, "隔离环"]],
    }

    def test_prefers_the_sheet_named_like_our_own_output(self):
        result = extract_names_from_sheets(self.TWO_SHEETS)
        assert result.sheet == PREFERRED_INPUT_SHEETS[0]
        assert list(result.names) == ["座板", "臂杆"]

    def test_still_reads_the_legacy_sheet_name(self):
        """改名兼容：2026-09-29 之前生成的清单，工作表还叫「补料清单」。

        只认新名的话，用户手上的旧文件会退化成「取第一张表」——
        多表工作簿（旧清单 + 备料参考）就会取错表。这条钉住兼容行为。
        """
        legacy = {
            LEGACY_SHEET_TITLE: [["序号", "名称"], [1, "座板"]],
            "备料参考": [["序号", "名称"], [1, "NECK"]],
        }
        result = extract_names_from_sheets(legacy)
        assert result.sheet == LEGACY_SHEET_TITLE
        assert list(result.names) == ["座板"]
        assert "NECK" not in result.names

    def test_new_name_wins_when_both_present(self):
        """新旧名同时出现时，认新的那张（万一有人手工留了一份旧表）。"""
        both = {
            SHEET_TITLE: [["序号", "名称"], [1, "新的"]],
            LEGACY_SHEET_TITLE: [["序号", "名称"], [1, "旧的"]],
        }
        result = extract_names_from_sheets(both)
        assert list(result.names) == ["新的"]

    def test_reference_sheet_no_longer_pollutes_the_list(self):
        """这条是那个 bug 的回归测试，别删。"""
        result = extract_names_from_sheets(self.TWO_SHEETS)
        for leaked in ("NECK", "HEAD_BACK", "隔离环"):
            assert leaked not in result.names, f"参考表的「{leaked}」又混进清单了"

    def test_ignored_sheets_are_reported(self):
        result = extract_names_from_sheets(self.TWO_SHEETS)
        assert result.skipped == ("备料参考",)
        assert result.ambiguous is True
        warning = result.warning()
        assert warning is not None
        assert "备料参考" in warning
        assert "--sheet" in warning

    def test_falls_back_to_a_worksheet_of_any_name(self):
        """没有「采购清单」这张表时，仍然要能读 —— 不能因为改名就罢工。"""
        result = extract_names_from_sheets({"Sheet1": [["名称"], ["垫片"]]})
        assert result.sheet == "Sheet1"
        assert list(result.names) == ["垫片"]

    def test_falls_back_to_the_first_candidate_not_the_union(self):
        sheets = {"甲": [["名称"], ["垫片"]], "乙": [["名称"], ["法兰盘"]]}
        result = extract_names_from_sheets(sheets)
        assert result.sheet == "甲"
        assert list(result.names) == ["垫片"]
        assert result.skipped == ("乙",)

    def test_sheets_without_a_name_column_are_not_candidates(self):
        sheets = {"采购清单": [["名称"], ["垫片"]], "说明": [["标题"], ["随便写点什么"]]}
        result = extract_names_from_sheets(sheets)
        assert result.sheet == "采购清单"
        assert result.skipped == ()
        assert result.ambiguous is False
        assert result.warning() is None

    def test_explicit_sheet_wins(self):
        result = extract_names_from_sheets(self.TWO_SHEETS, sheet="备料参考")
        assert result.sheet == "备料参考"
        assert "NECK" in result.names
        assert "座板" not in result.names

    def test_explicit_sheet_is_recorded_and_not_second_guessed(self):
        """用户自己指定了表，就别再提醒他「可以指定表」。"""
        result = extract_names_from_sheets(self.TWO_SHEETS, sheet="备料参考")
        assert result.requested is True
        assert result.skipped == (PREFERRED_INPUT_SHEETS[0],)
        assert result.warning() is None

    def test_auto_pick_is_not_marked_requested(self):
        assert extract_names_from_sheets(self.TWO_SHEETS).requested is False

    def test_explicit_sheet_that_does_not_exist(self):
        with pytest.raises(InputError, match="没有这张工作表"):
            extract_names_from_sheets(self.TWO_SHEETS, sheet="压根没有这张表")

    def test_explicit_sheet_without_a_name_column(self):
        sheets = {"采购清单": [["名称"], ["垫片"]], "说明": [["标题"], ["1"]]}
        result = extract_names_from_sheets(sheets, sheet="说明")
        assert result.names == ()
        assert result.sheet == "说明"

    def test_explicit_sheet_error_lists_what_is_available(self):
        with pytest.raises(InputError) as excinfo:
            extract_names_from_sheets(self.TWO_SHEETS, sheet="错的")
        assert "采购清单" in str(excinfo.value)  # 报错要告诉人有哪些表可选

    def test_single_sheet_needs_no_warning(self):
        result = extract_names_from_sheets({"采购清单": [["名称"], ["垫片"]]})
        assert result.ambiguous is False
        assert result.warning() is None


class TestExtractNames:
    def test_reads_from_file(self, tmp_path):
        path = _write_book(tmp_path / "a.xlsx", [["序号", "名称"], [1, "垫片"]])
        result = extract_names(path)
        assert list(result.names) == ["垫片"]
        assert "采购清单" in result.summary()

    def test_raises_when_no_name_column(self, tmp_path):
        path = _write_book(tmp_path / "a.xlsx", [["序号", "图号"], [1, "X-01"]])
        with pytest.raises(InputError, match=NAME_HEADER_KEYWORD):
            extract_names(path)

    def test_multi_sheet_file_is_no_longer_a_trap(self, tmp_path):
        """端到端回归：从真实文件读，也只认一张表。"""
        path = _write_book(
            tmp_path / "多表.xlsx",
            [["名称"], ["座板"]],
            extra_sheets={"备料参考": [["名称"], ["NECK"], ["HEAD_BACK"]]},
        )
        result = extract_names(path)
        assert list(result.names) == ["座板"]
        assert result.ambiguous is True
