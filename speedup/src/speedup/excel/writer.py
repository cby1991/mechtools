"""生成补料清单工作簿。

模板是**内置**的：程序自己画表头、边框、列宽、行高，不再依赖外部
``补料清单模板.xlsx``。这样脚本搬到任何机器上都能直接跑，也方便单测。

逐层拆开：
* :func:`build_supplement_workbook` —— 画空表（表头 + 格式行）
* :func:`fill_supplement_sheet` —— 把零件填进去，返回实际写入行数
* :func:`generate_supplement_list` —— 落盘成 ``<文件夹名>.xlsx``
"""

from __future__ import annotations

import os
from collections.abc import Sequence

from ..config import (
    BODY_FONT_NAME,
    COL_MATERIAL,
    COL_NAME,
    COL_NO,
    COL_NOTE,
    COL_QTY,
    DATA_ROW_HEIGHT,
    DATA_START_ROW,
    HEADER_FILL,
    HEADER_FONT_COLOR,
    HEADER_FONT_NAME,
    HEADER_ROW,
    HEADER_ROW_HEIGHT,
    MAX_ROWS,
    MIN_DATA_ROWS,
    SHEET_TITLE,
    SUPPLEMENT_HEADERS,
    SUPPLEMENT_WIDTHS,
)
from ..models import SupplementEntry

__all__ = [
    "build_supplement_workbook",
    "fill_supplement_sheet",
    "generate_supplement_list",
]


def _last_formatted_row(entry_count: int) -> int:
    """需要画格式的最后一行：至少撑到 ``MIN_DATA_ROWS``，且不超过 ``MAX_ROWS``。"""
    wanted = max(DATA_START_ROW + entry_count - 1, MIN_DATA_ROWS)
    return min(wanted, MAX_ROWS)


def build_supplement_workbook(entry_count: int = 0) -> tuple[object, object]:
    """建一个带完整格式的空补料清单。

    :param entry_count: 预计写入的零件数，用于决定要预画多少行格式
    :returns: ``(Workbook, Worksheet)``（openpyxl 对象，类型用 ``object`` 标注以避免硬依赖）
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_TITLE

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name=HEADER_FONT_NAME, size=11, bold=True, color=HEADER_FONT_COLOR)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_align = Alignment(horizontal="center", vertical="center")
    body_font = Font(name=BODY_FONT_NAME, size=11)
    body_align = Alignment(vertical="center")

    for index, title in enumerate(SUPPLEMENT_HEADERS, 1):
        cell = worksheet.cell(row=HEADER_ROW, column=index, value=title)
        cell.font, cell.fill, cell.alignment, cell.border = (
            header_font,
            header_fill,
            header_align,
            border,
        )
    worksheet.row_dimensions[HEADER_ROW].height = HEADER_ROW_HEIGHT

    for index, width in enumerate(SUPPLEMENT_WIDTHS, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    for row in range(DATA_START_ROW, _last_formatted_row(entry_count) + 1):
        worksheet.row_dimensions[row].height = DATA_ROW_HEIGHT
        for column in range(1, len(SUPPLEMENT_HEADERS) + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.font, cell.border, cell.alignment = body_font, border, body_align

    return workbook, worksheet


def fill_supplement_sheet(
    worksheet: object,
    entries: Sequence[SupplementEntry],
) -> int:
    """把零件写进工作表，返回实际写入行数（超出 ``MAX_ROWS`` 的部分丢弃）。"""
    written = 0
    for offset, entry in enumerate(entries):
        row = DATA_START_ROW + offset
        if row > MAX_ROWS:
            break
        worksheet.cell(row=row, column=COL_NO, value=offset + 1)  # type: ignore[attr-defined]
        worksheet.cell(row=row, column=COL_NAME, value=entry.name)  # type: ignore[attr-defined]
        worksheet.cell(row=row, column=COL_QTY, value=entry.qty)  # type: ignore[attr-defined]
        if entry.material:
            worksheet.cell(row=row, column=COL_MATERIAL, value=entry.material)  # type: ignore[attr-defined]
        if entry.note:
            worksheet.cell(row=row, column=COL_NOTE, value=entry.note)  # type: ignore[attr-defined]
        written += 1
    return written


def generate_supplement_list(
    folder: str | os.PathLike[str],
    entries: Sequence[SupplementEntry],
    filename: str | None = None,
) -> tuple[str, bool]:
    """生成 ``<文件夹名>.xlsx``（或指定文件名）并落盘。

    :returns: ``(保存路径, 是否覆盖了已存在的旧文件)``
    """
    target_dir = os.fspath(folder)
    name = filename or (os.path.basename(os.path.normpath(target_dir)) + ".xlsx")
    destination = os.path.join(target_dir, name)
    existed = os.path.exists(destination)

    workbook, worksheet = build_supplement_workbook(len(entries))
    fill_supplement_sheet(worksheet, entries)
    workbook.save(destination)  # type: ignore[attr-defined]
    return destination, existed
